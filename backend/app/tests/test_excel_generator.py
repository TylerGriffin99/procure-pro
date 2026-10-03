import os
from decimal import Decimal

from openpyxl import load_workbook
import io

from app.utils.excel_generator import generate_payment_recommendation_excel


def _sample_data():
    """Same structure as the PDF data dict from assessment_service.py."""
    return {
        "project_name": "Gilmours Central Seismic Upgrade",
        "project_number": "S-02314",
        "pr_number": 1,
        "revision": 0,
        "issue_date": "09 September 2025",
        "principal": "Foodstuffs North Island",
        "end_client_name": "Foodstuffs North Island",
        "end_client_address_lines": ["123 Main St", "Auckland"],
        "end_client_representative_first_name": "Owen",
        "landlord_split_pct": None,
        "operator_split_pct": None,
        "landlord_amount": Decimal("0"),
        "operator_amount": Decimal("0"),
        "engineer": "Owen Sanders",
        "contractor": "Kynoch Construction",
        "claim_number": 1,
        "claim_received": "29 August 2025",
        "payment_due": "20 September 2025",
        "contract_sum": Decimal("4172492.92"),
        "adjustment_to_provisional_sums": Decimal("0"),
        "approved_variation_orders": Decimal("0"),
        "adjusted_contract_sum": Decimal("4172492.92"),
        "value_claimed": Decimal("70317.39"),
        "adjustments": Decimal("-1730.00"),
        "total_recommended": Decimal("68587.39"),
        "retention_details": [
            {"percentage": "10%", "base": Decimal("68587.39"), "amount": Decimal("6858.74")},
        ],
        "total_retention": Decimal("6858.74"),
        "total_payment_to_date": Decimal("61728.65"),
        "previously_certified": Decimal("0"),
        "recommended_this_period": Decimal("61728.65"),
        "gst_rate": Decimal("0.15"),
        "gst_amount": Decimal("9259.30"),
        "total_including_gst": Decimal("70987.95"),
        "contract_works": [
            {
                "description": "Preliminaries",
                "contract_sum": Decimal("250000.00"),
                "contractor_claim": Decimal("6250.00"),
                "recommended": Decimal("6250.00"),
                "percentage": "3%",
                "variance": Decimal("0"),
                "previously_paid": Decimal("0"),
                "recommended_this_period": Decimal("6250.00"),
                "comments": "As per progress noted on site",
            },
        ],
        "variation_works": [
            {
                "ci_number": "CI-001",
                "description": "Additional drainage",
                "submission": Decimal("5000.00"),
                "type": "",
                "claimed_to_date": Decimal("5000.00"),
                "recommended": Decimal("4500.00"),
                "previously_paid": Decimal("0"),
                "this_period": Decimal("4500.00"),
                "status": "approved",
            },
        ],
        "provisional_sums": [
            {
                "ps_number": "PS-001",
                "description": "Electrical allowance",
                "contract_sum": Decimal("50000.00"),
                "claimed_to_date": Decimal("10000.00"),
                "recommended": Decimal("9500.00"),
                "previously_paid": Decimal("0"),
                "this_period": Decimal("9500.00"),
                "percentage": "19.0%",
                "status": "open",
            },
        ],
        "is_draft": True,
    }


def test_generate_excel_returns_valid_xlsx():
    """Generate an Excel file and verify it's a valid .xlsx workbook."""
    data = _sample_data()
    excel_bytes = generate_payment_recommendation_excel(data)

    assert excel_bytes[:2] == b"PK"
    assert len(excel_bytes) > 500

    wb = load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "Summary" in sheet_names
    assert "Contract Works" in sheet_names
    assert "Variations" in sheet_names
    assert "Provisional Sums" in sheet_names


def test_generate_excel_summary_sheet_has_key_values():
    """The Summary sheet should contain key financial data."""
    data = _sample_data()
    excel_bytes = generate_payment_recommendation_excel(data)
    wb = load_workbook(io.BytesIO(excel_bytes))
    ws = wb["Summary"]

    all_values = []
    for row in ws.iter_rows(values_only=True):
        all_values.extend([v for v in row if v is not None])

    assert any("Gilmours" in str(v) for v in all_values)
    assert any(abs(float(v) - 4172492.92) < 0.01 for v in all_values if isinstance(v, (int, float)))


def test_generate_excel_contract_works_rows():
    """Contract Works sheet should have a header row and data rows."""
    data = _sample_data()
    excel_bytes = generate_payment_recommendation_excel(data)
    wb = load_workbook(io.BytesIO(excel_bytes))
    ws = wb["Contract Works"]

    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) >= 2
    assert rows[1][0] == "Preliminaries"


def test_generate_excel_no_provisional_sums():
    """When provisional_sums is empty, sheet should still exist but have only headers."""
    data = _sample_data()
    data["provisional_sums"] = []
    excel_bytes = generate_payment_recommendation_excel(data)
    wb = load_workbook(io.BytesIO(excel_bytes))
    ws = wb["Provisional Sums"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 1


def test_generate_excel_writes_to_disk():
    """Write to /tmp for visual inspection."""
    data = _sample_data()
    excel_bytes = generate_payment_recommendation_excel(data)
    out_path = "/tmp/test_pr.xlsx"
    with open(out_path, "wb") as f:
        f.write(excel_bytes)
    assert os.path.exists(out_path)
