"""Generate payment recommendation Excel workbook using openpyxl."""

import io
from collections.abc import Sequence
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.payment_recommendation import PaymentRecommendation

NUMBER_FMT = "#,##0.00"
Cell = str | int | Decimal | None


def style_header_row(ws: Worksheet, row_num: int, col_count: int) -> None:
    header_font = Font(bold=True, size=9, color="FFFFFF")
    header_fill = PatternFill(start_color="1A2744", end_color="1A2744", fill_type="solid")
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_column_widths(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_table(ws: Worksheet, headers: list[str], rows: Sequence[Sequence[Cell]], widths: list[int]) -> None:
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    style_header_row(ws, 1, len(headers))
    for i, row in enumerate(rows, 2):
        for col, value in enumerate(row, 1):
            if isinstance(value, Decimal):
                ws.cell(row=i, column=col, value=float(value)).number_format = NUMBER_FMT
            else:
                ws.cell(row=i, column=col, value=value)
    set_column_widths(ws, widths)


def write_summary_sheet(wb: Workbook, rec: PaymentRecommendation) -> None:
    ws = wb.active
    ws.title = "Summary"

    title_font = Font(bold=True, size=14, color="1A2744")
    section_font = Font(bold=True, size=10, color="1A2744")
    label_font = Font(bold=True, size=9)
    value_font = Font(size=9)
    thin_border = Border(bottom=Side(style="thin", color="CCCCCC"))

    ws.merge_cells("A1:D1")
    ws["A1"] = f"Payment Recommendation #{rec.pr_number}"
    ws["A1"].font = title_font

    details = [
        ("Project", rec.project_name),
        ("Project Number", rec.project_number),
        ("Principal", rec.principal),
        ("Contractor", rec.contractor),
        ("Engineer", rec.engineer),
        ("Claim Number", rec.claim_number),
        ("Issue Date", rec.issue_date),
        ("Claim Received", rec.claim_received),
        ("Payment Due", rec.payment_due),
    ]
    financials = [
        ("Contract Sum", rec.contract_sum),
        ("Adjustment to Provisional Sums", rec.adjustment_to_provisional_sums),
        ("Approved Variation Orders", rec.approved_variation_orders),
        ("Adjusted Contract Sum", rec.adjusted_contract_sum),
        ("Value of Works Claimed to Date", rec.value_claimed),
        ("Adjustments", rec.adjustments),
        ("Total Value Recommended", rec.total_recommended),
        ("Total Retention", rec.total_retention),
        ("Total Payment to Date", rec.total_payment_to_date),
        ("Previously Certified", rec.previously_certified),
        ("Recommended This Period", rec.recommended_this_period),
        ("GST Amount", rec.gst_amount),
        ("Total Including GST", rec.total_including_gst),
    ]

    row = 3
    for label, value in details:
        ws.cell(row=row, column=1, value=label).font = label_font
        ws.cell(row=row, column=2, value=value).font = value_font
        ws.cell(row=row, column=1).border = thin_border
        ws.cell(row=row, column=2).border = thin_border
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Financial Summary").font = section_font
    row += 1
    for label, amount in financials:
        ws.cell(row=row, column=1, value=label).font = label_font
        cell = ws.cell(row=row, column=2, value=float(amount))
        cell.font = value_font
        cell.number_format = NUMBER_FMT
        ws.cell(row=row, column=1).border = thin_border
        cell.border = thin_border
        row += 1

    if rec.retention_details:
        row += 1
        ws.cell(row=row, column=1, value="Retention Breakdown").font = section_font
        row += 1
        for tier in rec.retention_details:
            ws.cell(row=row, column=1, value=tier.percentage).font = value_font
            cell = ws.cell(row=row, column=2, value=float(tier.amount))
            cell.font = value_font
            cell.number_format = NUMBER_FMT
            row += 1

    set_column_widths(ws, [30, 20, 15, 15])


def write_contract_works_sheet(wb: Workbook, rec: PaymentRecommendation) -> None:
    headers = [
        "Description",
        "Contract Sum",
        "Claim to Date",
        "Total Recommended",
        "Percentage",
        "Variance",
        "Previously Paid",
        "Recommended This Period",
        "Comments",
    ]
    rows = [
        (
            item.description,
            item.contract_sum,
            item.contractor_claim,
            item.recommended,
            item.percentage,
            item.variance,
            item.previously_paid,
            item.recommended_this_period,
            item.comments,
        )
        for item in rec.contract_works
    ]
    write_table(wb.create_sheet("Contract Works"), headers, rows, [30, 15, 15, 18, 12, 15, 15, 20, 30])


def write_variations_sheet(wb: Workbook, rec: PaymentRecommendation) -> None:
    headers = [
        "CI Number",
        "Description",
        "Submission",
        "Type",
        "Claimed to Date",
        "Recommended",
        "Previously Paid",
        "This Period",
        "Status",
        "Comments",
    ]
    rows = [
        (
            item.ci_number,
            item.description,
            item.submission,
            item.type,
            item.claimed_to_date,
            item.recommended,
            item.previously_paid,
            item.this_period,
            item.status,
            item.comments,
        )
        for item in rec.variation_works
    ]
    write_table(wb.create_sheet("Variations"), headers, rows, [15, 30, 15, 12, 15, 15, 15, 15, 12, 30])


def write_provisional_sums_sheet(wb: Workbook, rec: PaymentRecommendation) -> None:
    headers = [
        "PS Number",
        "Description",
        "Contract Sum",
        "Claimed to Date",
        "Recommended",
        "Previously Paid",
        "This Period",
        "Percentage",
        "Status",
        "Comments",
    ]
    rows = [
        (
            item.ps_number,
            item.description,
            item.contract_sum,
            item.claimed_to_date,
            item.recommended,
            item.previously_paid,
            item.this_period,
            item.percentage,
            item.status,
            item.comments,
        )
        for item in rec.provisional_sums
    ]
    write_table(wb.create_sheet("Provisional Sums"), headers, rows, [15, 30, 15, 15, 15, 15, 15, 12, 12, 30])


def generate_payment_recommendation_excel(rec: PaymentRecommendation) -> bytes:
    wb = Workbook()
    write_summary_sheet(wb, rec)
    write_contract_works_sheet(wb, rec)
    write_variations_sheet(wb, rec)
    write_provisional_sums_sheet(wb, rec)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
